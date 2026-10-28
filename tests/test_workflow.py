from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def test_exam_persists_and_updates(client, exam):
    created = client.post("/api/exams", json=exam)
    assert created.status_code == 201
    item_id = created.json()["id"]
    exam["title"] = "Revised title"
    assert client.put(f"/api/exams/{item_id}", json=exam).status_code == 200
    records = client.get("/api/exams").json()
    assert len(records) == 1
    assert records[0]["title"] == "Revised title"
    assert records[0]["tasks"][2]["criteria"][0]["points"] == 14
    assert client.put("/api/exams/unknown", json=exam).status_code == 404


def test_incorrect_points_rejected(client, exam):
    exam["tasks"][0]["criteria"][0]["points"] = -1
    assert client.post("/api/exams", json=exam).status_code == 422
    assert client.get("/api/exams").json() == []


def test_review_reports_real_issues_and_totals(client, exam):
    result = client.post("/api/review", json=exam).json()
    assert result["content_points"] == 40
    assert result["language_points"] == 60
    failed = [c["label"] for c in result["checks"] if not c["passed"]]
    assert failed == ["Formatprofil geprüft"]
    exam["tasks"][0]["criteria"][0]["evidence"] = "An invented quote"
    result = client.post("/api/review", json=exam).json()
    assert not next(c for c in result["checks"] if c["label"] == "Textbelege nachprüfbar")["passed"]


def test_profile_needs_reference_and_increments_version(client):
    profile = client.get("/api/profile").json()
    assert not profile["reviewed"]
    profile["reviewed"] = True
    assert client.put("/api/profile", json=profile).status_code == 422
    profile["source_reference"] = "Teacher-provided reference, page 2, year 2026"
    profile["font"] = "Times New Roman"
    profile["font_size"] = 12
    profile["margin_cm"] = 3
    saved = client.put("/api/profile", json=profile).json()
    assert saved["version"] == 2
    assert client.get("/api/profile").json() == saved


def test_word_exports_are_separate_and_use_saved_profile(client, exam):
    profile = client.get("/api/profile").json()
    profile.update(font="Times New Roman", font_size=12, margin_cm=3, line_spacing=2)
    assert client.put("/api/profile", json=profile).status_code == 200
    student = client.post("/api/export/exam", json=exam)
    rubric = client.post("/api/export/rubric", json=exam)
    assert student.status_code == rubric.status_code == 200
    assert "wordprocessingml" in student.headers["content-type"]
    doc = Document(BytesIO(student.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Outline the consequences." in text
    assert "Maya lost her job." in text
    assert "Benannt wird der Jobverlust." not in text
    assert not doc.tables
    assert len(doc.sections) == 2
    assert round(doc.sections[-1].left_margin.cm, 1) == 3
    assert round(doc.sections[-1].page_width.cm, 1) == 21
    assert doc.styles["Normal"].font.name == "Times New Roman"
    assert doc.styles["Normal"].font.size.pt == 12
    assert doc.styles["Normal"].paragraph_format.line_spacing == 2
    with ZipFile(BytesIO(student.content)) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        assert len(xml.xpath("//w:lnNumType", namespaces=NS)) == 1
        assert xml.xpath("//w:lnNumType/@w:countBy", namespaces=NS) == ["5"]
        assert xml.xpath("//w:suppressLineNumbers", namespaces=NS)
        styles = etree.fromstring(archive.read("word/styles.xml"))
        assert not styles.xpath('//w:style[@w:styleId="Title"]/w:pPr/w:pBdr', namespaces=NS)
        assert not styles.xpath('//w:style[@w:styleId="Title"]/w:rPr/w:rFonts/@w:asciiTheme', namespaces=NS)
        assert "Entwurf" in archive.read("word/footer1.xml").decode()
        assert 'w:instr="PAGE"' in archive.read("word/footer1.xml").decode()
    teacher = Document(BytesIO(rubric.content))
    assert len(teacher.tables) == 4
    assert teacher.tables[0].cell(1, 0).text == "Benannt wird der Jobverlust."
    assert teacher.tables[0].cell(1, 1).text == "Maya lost her job."
    assert teacher.tables[0].cell(1, 2).text == "10"
    assert teacher.tables[0].cell(1, 3).text == ""
    assert teacher.tables[0].columns[0].width > teacher.tables[0].columns[2].width * 4
    assert teacher.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert any("Gesamt: 100 Punkte" in p.text for p in teacher.paragraphs)


def test_export_requires_actual_material_and_rubric(client, exam):
    exam["tasks"][0]["criteria"] = []
    assert client.post("/api/export/rubric", json=exam).status_code == 422
    exam["source_text"] = ""
    assert client.post("/api/export/exam", json=exam).status_code == 422


def test_docx_import_preserves_criteria_in_tables(client):
    doc = Document()
    doc.add_paragraph("Example rubric")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Analyse narrative perspective"
    table.cell(0, 1).text = "8 points"
    output = BytesIO()
    doc.save(output)
    imported = client.post("/api/import", files={"file": ("rubric.docx", output.getvalue())})
    assert imported.status_code == 200
    assert "Analyse narrative perspective | 8 points" in imported.json()["text"]
    assert client.get("/api/knowledge").json() == []
    # Import does not silently add or approve knowledge. The teacher reviews first.
    item = {"title": "Analysis example", "kind": "Beispiel", "content": imported.json()["text"], "reference": "Original teacher example", "reviewed": True}
    created = client.post("/api/knowledge", json=item).json()
    assert client.get("/api/knowledge").json()[0]["content"] == item["content"]
    item.update(content="A corrected criterion", reviewed=False)
    updated = client.put(f'/api/knowledge/{created["id"]}', json=item)
    assert updated.status_code == 200
    assert not updated.json()["reviewed"]
    assert client.delete(f'/api/knowledge/{created["id"]}').status_code == 204
    assert client.get("/api/knowledge").json() == []


def test_pdf_import_extracts_real_text(client):
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (A reusable knowledge source.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    response = client.post("/api/import", files={"file": ("source.pdf", output.getvalue())})
    assert response.status_code == 200
    assert response.json()["text"] == "A reusable knowledge source."


def test_import_failure_modes_are_actionable(client):
    assert client.post("/api/import", files={"file": ("a.txt", b"\xef\xbb\xbfHello")}).json()["text"] == "Hello"
    assert client.post("/api/import", files={"file": ("a.exe", b"Hello")}).status_code == 415
    assert client.post("/api/import", files={"file": ("a.pdf", b"not a PDF")}).status_code == 422
    assert client.post("/api/import", files={"file": ("a.txt", b"A" * 60001)}).status_code == 422
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    response = client.post("/api/import", files={"file": ("scan.pdf", output.getvalue())})
    assert response.status_code == 422
    assert "OCR" in response.json()["detail"]


def test_unreferenced_knowledge_cannot_be_reviewed(client):
    item = {"title": "A rule", "content": "Use a known layout.", "reviewed": True}
    assert client.post("/api/knowledge", json=item).status_code == 422


def test_cross_site_writes_rejected(client, exam):
    assert client.post("/api/exams", json=exam, headers={"Origin": "https://other.example"}).status_code == 403
    assert client.post("/api/exams", json=exam, headers={"Origin": "http://testserver"}).status_code == 201
    assert client.get("/api/exams").headers["cache-control"] == "no-store"


def test_missing_ai_key_is_honest_and_manual_work_still_works(client, exam):
    assert not client.get("/api/status").json()["ai_available"]
    assert client.post("/api/generate", json={"exam": exam}).status_code == 503
    assert client.post("/api/exams", json=exam).status_code == 201
