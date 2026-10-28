from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .models import Exam, FormatProfile


def suppress_line_numbers(paragraph):
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:suppressLineNumbers"))


def make_table(doc, headings, proportions, width_cm):
    table = doc.add_table(rows=1, cols=len(headings))
    table.style = "Table Grid"
    table.autofit = False
    for column, cell, proportion, text in zip(table.columns, table.rows[0].cells, proportions, headings):
        column.width = cell.width = Cm(width_cm * proportion)
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def add_table_row(table, values):
    row = table.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for cell, value in zip(row.cells, values):
        cell.text = value


def make_document(exam: Exam, profile: FormatProfile, kind: str) -> BytesIO:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(profile.margin_cm)
    section.left_margin = section.right_margin = Cm(profile.margin_cm)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = profile.font
        style.font.size = Pt(profile.font_size)
        style.font.color.rgb = None
        style.paragraph_format.line_spacing = profile.line_spacing
        style.paragraph_format.space_after = Pt(6)
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        # Theme fonts otherwise override explicit face names in some Word renderers.
        for attribute in list(fonts.attrib):
            if attribute.lower().endswith("theme"):
                del fonts.attrib[attribute]
        fonts.set(qn("w:eastAsia"), profile.font)
        fonts.set(qn("w:cs"), profile.font)
    title_properties = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_properties.findall(qn("w:pBdr")):
        title_properties.remove(border)
    doc.styles["Title"].font.size = Pt(profile.font_size + 5)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.bold = True
    footer = section.footer.paragraphs[0]
    footer.add_run(f"Entwurf · Formatprofil v{profile.version} · Seite ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = exam.title
    doc.core_properties.subject = "Englischklausur – redaktioneller Entwurf"
    doc.core_properties.comments = f"Format: {profile.name}; Referenz: {profile.source_reference or 'nicht hinterlegt'}"
    doc.add_heading(exam.title, 0)
    doc.add_paragraph(f"Englisch · {exam.course} · {exam.duration} Minuten")
    if kind == "exam":
        doc.add_paragraph("Name: ____________________________    Datum: ______________")
        doc.add_paragraph(f"Hilfsmittel: {exam.aids}")
        doc.add_heading("Tasks", 1)
        for i, task in enumerate(exam.tasks, 1):
            doc.add_paragraph(f"{i}. {task.instruction}")
        for paragraph in doc.paragraphs:
            suppress_line_numbers(paragraph)
        source_section = doc.add_section(WD_SECTION.NEW_PAGE)
        suppress_line_numbers(doc.add_heading(exam.source_title or "Text", 1))
        suppress_line_numbers(doc.add_paragraph(exam.source_reference or "Quellenangabe ergänzen"))
        # Word numbers its rendered lines; no misleading precomputed browser wrapping.
        line_numbers = OxmlElement("w:lnNumType")
        line_numbers.set(qn("w:countBy"), "5")
        line_numbers.set(qn("w:restart"), "newPage")
        source_section._sectPr.append(line_numbers)
        for paragraph in exam.source_text.split("\n"):
            doc.add_paragraph(paragraph)
    else:
        doc.add_heading("Erwartungshorizont", 1)
        doc.add_paragraph("Redaktioneller Entwurf. Die Lehrkraft prüft fachliche Richtigkeit, alternative Lösungen und die zur Klausur passende Gewichtung.")
        total_content = 0
        for i, task in enumerate(exam.tasks, 1):
            total = sum(c.points for c in task.criteria)
            total_content += total
            doc.add_heading(f"Aufgabe {i} · AFB {task.area} · {total} Punkte", 2)
            doc.add_paragraph(task.instruction).paragraph_format.keep_with_next = True
            table = make_table(doc, ("Erwartete Leistung", "Textbeleg", "Max.", "Erreicht"), (.49, .29, .09, .13), 21 - 2 * profile.margin_cm)
            for criterion in task.criteria:
                add_table_row(table, (criterion.description, criterion.evidence, str(criterion.points), ""))
        doc.add_heading("Sprachliche Leistung / Darstellungsleistung", 2)
        table = make_table(doc, ("Kriterium", "Max.", "Erreicht"), (.76, .1, .14), 21 - 2 * profile.margin_cm)
        for criterion in exam.language_criteria:
            add_table_row(table, (criterion.description, str(criterion.points), ""))
        language = sum(c.points for c in exam.language_criteria)
        doc.add_paragraph(f"Inhalt: {total_content} · Sprache/Darstellung: {language} · Gesamt: {total_content + language} Punkte")
        doc.add_heading("Prüfvermerk", 2)
        doc.add_paragraph(f"Formatprofil: {profile.name} (v{profile.version})\nReferenz: {profile.source_reference or 'Noch nicht hinterlegt'}\nFormat manuell geprüft: {'Ja' if profile.reviewed else 'Nein'}")
        if exam.generation_notes:
            doc.add_paragraph(exam.generation_notes)
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
