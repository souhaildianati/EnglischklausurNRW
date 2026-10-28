import os
from io import BytesIO
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader

from . import ai, store
from .documents import make_document
from .models import Exam, FormatProfile, GenerateRequest, KnowledgeInput
from .pdf_analysis import analyze_pdf
from .review import review

app = FastAPI(title="Klausurwerk", version="0.1.0")


@app.middleware("http")
async def local_app_protection(request: Request, call_next):
    # No public accounts yet. Reject cross-site writes, including multipart uploads.
    if request.method not in ("GET", "HEAD", "OPTIONS"):
        origin = request.headers.get("origin")
        if origin:
            from urllib.parse import urlsplit
            if urlsplit(origin).netloc != request.headers.get("host"):
                return JSONResponse({"detail": "Anfrage von fremder Herkunft abgewiesen."}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
    return response


@app.get("/api/status")
def status():
    return {"status": "ok", "ai_available": bool(os.environ.get("LLM_API_KEY")), "ai_model": os.environ.get("LLM_MODEL", "gpt-4.1-mini"), "mode": "private-prototype"}


@app.get("/api/exams")
def list_exams():
    return store.all_items("exams")


@app.post("/api/exams", status_code=201)
def create_exam(exam: Exam):
    return store.save_item("exams", exam.model_dump())


@app.put("/api/exams/{exam_id}")
def update_exam(exam_id: str, exam: Exam):
    if not store.get_item("exams", exam_id):
        raise HTTPException(404, "Klausur nicht gefunden.")
    return store.save_item("exams", exam.model_dump(), exam_id)


@app.post("/api/review")
def review_exam(exam: Exam):
    return review(exam, store.profile())


@app.post("/api/export/{kind}")
def export_exam(kind: Literal["exam", "rubric"], exam: Exam):
    if not exam.source_text.strip() or not exam.tasks:
        raise HTTPException(422, "Für den Export sind ein Ausgangstext und mindestens eine Aufgabe erforderlich.")
    if kind == "rubric" and not all(t.criteria for t in exam.tasks):
        raise HTTPException(422, "Bitte für jede Aufgabe Bewertungskriterien ergänzen.")
    name = "Klausur" if kind == "exam" else "Erwartungshorizont"
    return StreamingResponse(make_document(exam, store.profile(), kind), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{name}.docx"'})


@app.get("/api/profile")
def get_profile():
    return store.profile()


@app.put("/api/profile")
def update_profile(profile: FormatProfile):
    return store.save_profile(profile)


@app.get("/api/knowledge")
def knowledge():
    return store.all_items("knowledge")


@app.post("/api/knowledge", status_code=201)
def add_knowledge(item: KnowledgeInput):
    return store.save_item("knowledge", item.model_dump())


@app.put("/api/knowledge/{item_id}")
def update_knowledge(item_id: str, item: KnowledgeInput):
    if not store.get_item("knowledge", item_id):
        raise HTTPException(404, "Wissenseintrag nicht gefunden.")
    return store.save_item("knowledge", item.model_dump(), item_id)


@app.delete("/api/knowledge/{item_id}", status_code=204)
def delete_knowledge(item_id: str):
    if not store.delete_item("knowledge", item_id):
        raise HTTPException(404, "Wissenseintrag nicht gefunden.")


@app.post("/api/import")
async def import_text(file: UploadFile):
    data = await file.read(10 * 1024 * 1024 + 1)
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "Dateien dürfen höchstens 10 MB groß sein.")
    suffix = Path(file.filename or "").suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(422, "Verschlüsselte PDFs werden nicht unterstützt.")
            if len(reader.pages) > 100:
                raise HTTPException(422, "Bitte einen Ausschnitt mit höchstens 100 Seiten importieren.")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            import zipfile
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                    raise HTTPException(413, "Die entpackte Word-Datei ist zu groß.")
            doc = Document(BytesIO(data))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        elif suffix in (".txt", ".md"):
            text = data.decode("utf-8-sig")
        else:
            raise HTTPException(415, "Unterstützt werden PDF, DOCX, TXT und Markdown.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Die Datei konnte nicht gelesen werden. Bitte als unverschlüsselte PDF, DOCX oder UTF-8-Textdatei speichern.") from None
    if not text.strip():
        raise HTTPException(422, "Kein lesbarer Text gefunden. Gescannte PDFs benötigen vorher eine Texterkennung (OCR).")
    if len(text) > 60000:
        raise HTTPException(422, "Der Text ist zu lang. Bitte einen Ausschnitt mit höchstens 60.000 Zeichen verwenden.")
    return {"text": text.strip(), "filename": file.filename, "notice": "Extraktion prüfen: Tabellen, Seitenumbrüche und PDF-Lesereihenfolge können abweichen."}


@app.post("/api/generate")
async def generate_draft(request: GenerateRequest):
    return await ai.generate(request.exam)


@app.post("/api/analyze-pdf")
def analyze_reference_pdf(file: UploadFile):
    if Path(file.filename or "").suffix.lower() != ".pdf":
        raise HTTPException(415, "Für die Layoutanalyse bitte die Original-PDF auswählen.")
    data = file.file.read(10 * 1024 * 1024 + 1)
    return analyze_pdf(data, Path(file.filename).name)


DIST = Path(__file__).resolve().parent.parent / "frontend" / "dist"
if DIST.exists():
    app.mount("/assets", StaticFiles(directory=DIST / "assets"), name="assets")

    @app.get("/")
    def index():
        return FileResponse(DIST / "index.html")
